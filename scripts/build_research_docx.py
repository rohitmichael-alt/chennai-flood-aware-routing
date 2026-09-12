"""Build the research manuscript DOCX from its authoritative Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PROJECT_RESEARCH_AND_EVIDENCE.md"
OUTPUT = ROOT / "docs" / "PROJECT_RESEARCH_AND_EVIDENCE.docx"
NAVY = RGBColor(31, 78, 121)
MUTED = RGBColor(90, 98, 108)
LIGHT = "EAF0F6"


def set_font(run, name="Calibri", size=11, *, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "").replace("\\(", "").replace("\\)", "")


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def configure_table(table):
    table.autofit = False
    columns = len(table.columns)
    widths = [9360 // columns] * columns
    widths[-1] += 9360 - sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shade(cell, LIGHT)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    set_font(run, size=8.5, bold=row_index == 0)


def page_field(paragraph):
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, before, after in (
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = NAVY
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208


def add_cover(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    header = section.header.paragraphs[0]
    header.text = "CHENNAI COMPOUND-DISRUPTION ROUTING | RESEARCH MANUSCRIPT"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(header.runs[0], size=8.5, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field(footer)
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(kicker.add_run("REPRODUCIBLE SYSTEMS AND EVALUATION REPORT"), size=10, bold=True, color=NAVY)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(14)
    set_font(
        title.add_run("Certificate-Gated Dynamic Routing\nunder Compound Urban Disruptions"),
        size=25,
        bold=True,
        color=RGBColor(26, 55, 82),
    )
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(
        subtitle.add_run("A Chennai-oriented integration and evidence package"),
        size=14,
        italic=True,
        color=MUTED,
    )
    doc.add_paragraph()
    status = doc.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(status.add_run("Execution status: Stage 10 partial, with explicit experiment gaps"), size=11, bold=True, color=NAVY)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(date.add_run("Evidence regenerated September 2026"), size=10, color=MUTED)
    doc.add_page_break()


def add_image(doc: Document, path: Path, caption: str):
    if not path.is_file():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = paragraph.add_run().add_picture(str(path), width=Inches(6.2))
    doc_pr = picture._inline.docPr
    doc_pr.set("descr", caption)
    doc_pr.set("title", path.stem)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    set_font(cap.add_run(caption), size=9, italic=True, color=MUTED)


def build():
    doc = Document()
    configure_styles(doc)
    add_cover(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    index = 1  # skip source H1; cover already supplies it
    while index < len(lines):
        line = lines[index]
        if line.strip() == "\\[":
            equation: list[str] = []
            index += 1
            while index < len(lines) and lines[index].strip() != "\\]":
                equation.append(lines[index].strip())
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            set_font(paragraph.add_run(" ".join(equation)), name="Cambria Math", size=10.5)
            index += 1
            continue
        if line.startswith("```"):
            if in_code:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.space_after = Pt(8)
                shade_proxy = OxmlElement("w:shd")
                shade_proxy.set(qn("w:fill"), "F4F6F9")
                paragraph._p.get_or_add_pPr().append(shade_proxy)
                set_font(paragraph.add_run("\n".join(code_lines)), name="Consolas", size=8.5)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[ :\-|]+\|$", lines[index + 1]):
            rows: list[list[str]] = []
            rows.append([clean_inline(cell.strip()) for cell in line.strip("|").split("|")])
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([clean_inline(cell.strip()) for cell in lines[index].strip("|").split("|")])
                index += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
            configure_table(table)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            level = min(len(heading.group(1)) - 1, 3)
            text = clean_inline(heading.group(2))
            doc.add_heading(text, level=level)
            if text.startswith("Stage 1 "):
                add_image(doc, ROOT / "outputs" / "maps" / "stage1_publication_before_after.png", "Figure 1. Dated Stage 3 route before and after the historical-hotspot scenario blockage.")
            if text.startswith("6.3 OpenStreetMap"):
                add_image(doc, ROOT / "docs" / "evidence" / "STAGE3_GRAPH_QA_MAP.png", "Figure 2. Stage 3 Chennai road-graph quality-assurance map.")
        elif re.match(r"^\d+\.\s+", line):
            doc.add_paragraph(clean_inline(re.sub(r"^\d+\.\s+", "", line)), style="List Number")
        elif line.startswith("- "):
            doc.add_paragraph(clean_inline(line[2:]), style="List Bullet")
        elif line.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.right_indent = Inches(0.25)
            paragraph.paragraph_format.space_after = Pt(10)
            set_font(paragraph.add_run(clean_inline(line[2:])), size=11, italic=True, color=RGBColor(55, 70, 85))
        elif line.strip():
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.add_run(clean_inline(line))
        index += 1
    doc.core_properties.title = "Certificate-Gated Dynamic Routing under Compound Urban Disruptions"
    doc.core_properties.subject = "Chennai-oriented reproducible routing integration and evaluation"
    doc.core_properties.author = "Research project team"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
